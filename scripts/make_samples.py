"""Generate one sample file per supported type into data/samples/ for a quick test."""
from pathlib import Path

import docx
import pandas as pd
import pymupdf
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

out = Path(__file__).resolve().parent.parent / "data" / "samples"
out.mkdir(parents=True, exist_ok=True)

# Text PDF
pdf = pymupdf.open()
page = pdf.new_page()
page.insert_text((72, 72), "Acme Corp Refund Policy", fontsize=16)
page.insert_textbox(pymupdf.Rect(72, 100, 520, 400),
    "Customers may request a full refund within 30 days of purchase. "
    "After 30 days, store credit is issued instead. Refunds are processed "
    "within 5 business days to the original payment method.", fontsize=11)
pdf.save(out / "refund_policy.pdf")


def text_image(lines, size=(1400, 400)):
    img = Image.new("RGB", size, "white")
    d = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 40)
    except OSError:
        font = ImageFont.load_default(size=40)
    for i, line in enumerate(lines):
        d.text((40, 40 + i * 70), line, fill="black", font=font)
    return img


# Scanned PDF (image only, no text layer) -> must be OCR'd
scan = text_image(["Warranty: all hardware is covered", "for two years against defects."])
scan.save(out / "_scan.png")
spdf = pymupdf.open()
spage = spdf.new_page(width=700, height=200)
spage.insert_image(spage.rect, filename=str(out / "_scan.png"))
spdf.save(out / "scanned_warranty.pdf")
(out / "_scan.png").unlink()

# Image -> OCR
text_image(["Support hours: Monday to Friday", "9am to 6pm Eastern Time."]).save(out / "support_hours.png")

# Word
d = docx.Document()
d.add_heading("Employee Onboarding", level=1)
d.add_paragraph("New employees receive a laptop on their first day and must complete security training within one week.")
t = d.add_table(rows=2, cols=2)
t.cell(0, 0).text, t.cell(0, 1).text = "Item", "Owner"
t.cell(1, 0).text, t.cell(1, 1).text = "Laptop setup", "IT Helpdesk"
d.save(out / "onboarding.docx")

# PowerPoint
prs = Presentation()
s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = "Q3 Roadmap"
s.placeholders[1].text = "Launch the mobile app in September and expand to Canada in Q4."
prs.save(out / "roadmap.pptx")

# Excel (two sheets)
with pd.ExcelWriter(out / "sales.xlsx") as w:
    pd.DataFrame({"Region": ["North", "South", "West"], "Revenue": [120000, 95000, 143000],
                  "Manager": ["Alice", "Bob", "Carla"]}).to_excel(w, sheet_name="Q3", index=False)
    pd.DataFrame({"Product": ["Widget", "Gadget"], "Units": [500, 320]}).to_excel(w, sheet_name="Products", index=False)

# CSV + text
pd.DataFrame({"Office": ["New York", "London"], "Staff": [40, 25]}).to_csv(out / "offices.csv", index=False)
(out / "faq.md").write_text("# FAQ\n\nShipping is free on orders over $50. International shipping takes 7-10 days.\n")

print("Samples written to", out)
for f in sorted(out.iterdir()):
    print("  ", f.name)
