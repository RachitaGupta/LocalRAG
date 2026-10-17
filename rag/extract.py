"""Turn files of many types into plain-text "sections" with metadata.

Each extractor yields dicts: {"text": str, "meta": {...}} where meta holds
location info (page, sheet, slide) used later for citations.
"""
from __future__ import annotations

import io
import logging
from pathlib import Path
from typing import Iterator

from . import config

log = logging.getLogger(__name__)

PDF_EXT = {".pdf"}
DOCX_EXT = {".docx"}
PPTX_EXT = {".pptx"}
EXCEL_EXT = {".xlsx", ".xlsm", ".xls"}
CSV_EXT = {".csv", ".tsv"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp", ".gif"}
TEXT_EXT = {".txt", ".md", ".markdown", ".json", ".html", ".htm", ".xml", ".log"}

SUPPORTED_EXT = PDF_EXT | DOCX_EXT | PPTX_EXT | EXCEL_EXT | CSV_EXT | IMAGE_EXT | TEXT_EXT

Section = dict


# ---------------------------------------------------------------- OCR helpers
def _ocr_image(image) -> str:
    """OCR a PIL image. Returns '' if OCR is disabled or Tesseract is missing."""
    if not config.OCR_ENABLED:
        return ""
    try:
        import pytesseract
        if config.TESSERACT_CMD:
            pytesseract.pytesseract.tesseract_cmd = config.TESSERACT_CMD
        return pytesseract.image_to_string(image, lang=config.OCR_LANG)
    except Exception as exc:  # TesseractNotFoundError, bad image, etc.
        log.warning("OCR failed: %s", exc)
        return ""


# ---------------------------------------------------------------- extractors
def extract_pdf(path: Path) -> Iterator[Section]:
    import pymupdf
    from PIL import Image

    with pymupdf.open(path) as doc:
        for i, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            method = "text"
            # Scanned page (little/no text layer) -> render and OCR it
            if len(text) < config.OCR_MIN_CHARS:
                pix = page.get_pixmap(dpi=config.OCR_DPI)
                img = Image.open(io.BytesIO(pix.tobytes("png")))
                ocr_text = _ocr_image(img).strip()
                if len(ocr_text) > len(text):
                    text, method = ocr_text, "ocr"
            if text:
                yield {"text": text, "meta": {"page": i, "method": method}}


def extract_docx(path: Path) -> Iterator[Section]:
    import docx

    d = docx.Document(str(path))
    parts: list[str] = []
    for p in d.paragraphs:
        if p.text.strip():
            # Keep headings recognisable so chunks carry structure
            if p.style is not None and p.style.name.lower().startswith("heading"):
                parts.append(f"\n## {p.text.strip()}\n")
            else:
                parts.append(p.text.strip())
    for t_idx, table in enumerate(d.tables, start=1):
        rows = [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        if rows:
            parts.append(f"\n[Table {t_idx}]\n" + "\n".join(rows))
    text = "\n".join(parts).strip()
    if text:
        yield {"text": text, "meta": {}}

    # OCR images embedded in the document
    if config.OCR_ENABLED:
        from PIL import Image
        for rel in d.part.rels.values():
            if "image" in rel.reltype:
                try:
                    img = Image.open(io.BytesIO(rel.target_part.blob))
                    ocr = _ocr_image(img).strip()
                    if ocr:
                        yield {"text": ocr, "meta": {"method": "ocr-embedded-image"}}
                except Exception as exc:
                    log.debug("Skipping embedded image: %s", exc)


def extract_pptx(path: Path) -> Iterator[Section]:
    from pptx import Presentation

    prs = Presentation(str(path))
    for i, slide in enumerate(prs.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            parts.append("Notes: " + slide.notes_slide.notes_text_frame.text.strip())
        if parts:
            yield {"text": "\n".join(parts), "meta": {"slide": i}}


def _frame_blocks(df, meta: dict) -> Iterator[Section]:
    """Render a DataFrame as row-wise 'col: value' text in blocks of N rows.

    Row-wise text embeds far better than raw grids, and repeating the
    header names in every block keeps each chunk self-describing.
    """
    df = df.dropna(how="all").dropna(axis=1, how="all")
    if df.empty:
        return
    df.columns = [str(c).strip() for c in df.columns]
    n = config.EXCEL_ROWS_PER_BLOCK
    for start in range(0, len(df), n):
        block = df.iloc[start:start + n]
        lines = []
        for idx, row in block.iterrows():
            cells = [f"{col}: {val}" for col, val in row.items() if str(val).strip() not in ("", "nan")]
            if cells:
                lines.append("; ".join(cells))
        if lines:
            yield {
                "text": "\n".join(lines),
                "meta": {**meta, "rows": f"{start + 2}-{start + len(block) + 1}"},
            }


def extract_excel(path: Path) -> Iterator[Section]:
    import pandas as pd

    sheets = pd.read_excel(path, sheet_name=None)  # dict of all sheets
    for name, df in sheets.items():
        yield from _frame_blocks(df, {"sheet": str(name)})


def extract_csv(path: Path) -> Iterator[Section]:
    import pandas as pd

    sep = "\t" if path.suffix.lower() == ".tsv" else None
    df = pd.read_csv(path, sep=sep, engine="python")
    yield from _frame_blocks(df, {})


def extract_image(path: Path) -> Iterator[Section]:
    from PIL import Image

    with Image.open(path) as img:
        img.load()
        text = _ocr_image(img.convert("RGB")).strip()
    if text:
        yield {"text": text, "meta": {"method": "ocr"}}


def extract_text(path: Path) -> Iterator[Section]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    if path.suffix.lower() in {".html", ".htm"}:
        import re
        text = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", text, flags=re.S | re.I)
        text = re.sub(r"<[^>]+>", " ", text)
    if text.strip():
        yield {"text": text, "meta": {}}


# ---------------------------------------------------------------- dispatcher
def extract(path: str | Path) -> list[Section]:
    path = Path(path)
    ext = path.suffix.lower()
    if ext in PDF_EXT:
        fn = extract_pdf
    elif ext in DOCX_EXT:
        fn = extract_docx
    elif ext in PPTX_EXT:
        fn = extract_pptx
    elif ext in EXCEL_EXT:
        fn = extract_excel
    elif ext in CSV_EXT:
        fn = extract_csv
    elif ext in IMAGE_EXT:
        fn = extract_image
    elif ext in TEXT_EXT:
        fn = extract_text
    else:
        raise ValueError(f"Unsupported file type: {ext}")

    sections = list(fn(path))
    for s in sections:
        s["meta"] = {"source": str(path), "filename": path.name, "filetype": ext.lstrip("."), **s["meta"]}
    return sections
